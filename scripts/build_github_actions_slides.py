"""Generate the GitHub Actions beginner workshop deck: python -m pip install python-pptx==1.0.2."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333), Inches(7.5)
BG, PANEL, INK, MUTED, ACCENT = '101D30', '1B2C43', 'F4F7FB', 'B5C4D9', 'BBACFF'

def box(s,x,y,w,h,color):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
    return sh

def text(s,x,y,w,h,content,size=23,color=INK,font='Aptos',bold=False):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02); tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12 if font!='Consolas' else 3)
        if font=='Consolas': p.line_spacing=1.0
    return sh

def slide(section,title,sub=None):
    s=r.slides.add_slide(r.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    box(s,0,0,.16,7.5,ACCENT)
    text(s,.6,.35,12,.3,section.upper(),12,ACCENT,bold=True)
    text(s,.6,.88,12.1,.65,title,32,bold=True)
    if sub: text(s,.6,1.62,12,.5,sub,17,MUTED)
    box(s,.6,6.98,12.1,.012,'38516D')
    text(s,.6,7.13,11,.22,'DISTRITEK DEVOPS   /   Muhammad Asdar   /   GitHub Actions',10,MUTED)
    text(s,12.05,7.08,.6,.3,str(len(r.slides)).zfill(2),12,ACCENT)
    return s

def notes(s,body): s.notes_slide.notes_text_frame.text=body

def cards(section,title,items,sub=None,note=''):
    s=slide(section,title,sub); n=len(items); w=(12.1-.25*(n-1))/n
    for i,(head,body) in enumerate(items):
        x=.6+i*(w+.25); box(s,x,2.3,w,3.9,PANEL)
        text(s,x+.22,2.57,w-.44,.8,head,24,ACCENT,bold=True)
        text(s,x+.22,3.5,w-.44,2.5,body,21)
    notes(s,note); return s

def code(section,title,body,explanation,sub=None,note=''):
    s=slide(section,title,sub); box(s,.6,2.25,8,4.38,PANEL)
    longest=max(map(len,body.splitlines())); count=len(body.splitlines())
    size=min(20, 890/max(longest,1), 278/max(count,1)-3)
    text(s,.85,2.47,7.5,3.95,body,size,INK,'Consolas')
    text(s,8.98,2.4,3.7,3.9,explanation,22)
    notes(s,note); return s

s=slide('Workshop pemula • 120 menit','GitHub Actions: otomasi dari commit pertama')
text(s,.7,2.25,8.1,1.6,'Biar pemeriksaan rutin\ndikerjakan otomatis.',34,bold=True)
text(s,.7,4.25,8,1.1,'Kenali CI/CD, tulis workflow,\nbaca hasil, dan perbaiki kegagalan.',24,MUTED)
box(s,9.3,2.3,3.2,3.6,PANEL)
text(s,9.6,2.65,2.6,2.9,'COMMIT\n↓\nCHECK\n↓\nRESULT',27,ACCENT,bold=True)
text(s,.7,6.1,10,.4,'Muhammad Asdar  ·  Distritek DevOps',19)
notes(s,'Pengantar untuk peserta yang baru mengenal GitHub Actions. Asumsi: peserta mengenal file, terminal dasar, commit, dan repository. Bahan pendamping: praktik-github-actions.md. GitHub Actions dapat digunakan tanpa Docker.')
cards('Pengantar','Setiap perubahan kode perlu diperiksa',[('Pekerjaan berulang','Ambil kode terbaru.\nSiapkan versi runtime.\nJalankan test dan build.'),('Tantangannya','Langkah bisa terlupa.\nHasil sulit ditelusuri.\nError baru terlihat setelah aplikasi dipakai.')],note='Pembuka: Setelah belajar menjalankan aplikasi, bayangkan project dikerjakan lima orang. Setiap perubahan perlu diperiksa. Apakah kita mau mengingat dan mengetik semua langkah itu setiap kali?')
cards('Pengantar','GitHub Actions menjalankan instruksi kita',[('Kita menulis','Urutan pekerjaan di file YAML.\nTentukan kapan dijalankan dan apa yang diperiksa.'),('Runner mengerjakan','Mesin mengeksekusi instruksi.\nGitHub menampilkan log, status, dan hasilnya.')],sub='Otomasi tidak menebak test yang dibutuhkan; kita yang mendefinisikannya.',note='Analogi: workflow adalah daftar tugas, event adalah tanda mulai, runner adalah pelaksana. Mulai dari contoh mencetak pesan sebelum membahas deployment. https://docs.github.com/en/actions/get-started/understand-github-actions')
cards('Pengantar','Git, GitHub, Actions, dan Docker',[('Git + GitHub','Git mencatat riwayat kode.\nGitHub menyimpan repo dan membantu kolaborasi.'),('GitHub Actions','Mengotomatiskan pekerjaan saat terjadi event pada repo.'),('Docker','Mengemas dan menjalankan aplikasi.\nDapat digunakan di workflow, tetapi tidak wajib.')],note='Jangan menyamakan Actions dengan hosting aplikasi permanen. Runner CI adalah tempat kerja sementara, bukan server production.')
cards('CI/CD','Apa yang dimaksud CI dan CD?',[('CI','Continuous Integration.\nPerubahan kode diperiksa lewat test dan build otomatis.'),('Delivery','Hasil yang lulus disiapkan untuk rilis.\nDeploy dapat menunggu keputusan manusia.'),('Deployment','Perubahan yang lulus otomatis diterapkan ke environment tujuan.')],sub='CI hijau berarti pemeriksaan yang ditulis lulus; bukan jaminan bebas bug.',note='Jelaskan singkatan CD bisa merujuk delivery atau deployment. Dalam workshop ini CI dipraktikkan; deployment dibahas sebagai tahap lanjutan.')
cards('Tujuan & rute','Dari workflow pertama ke CI yang berguna',[('25 menit','Pengantar dan istilah.\nPersiapan repository latihan.'),('55 menit','Lab 1 Hello: 15 menit.\nLab 2 CI Node: 25 menit.\nLab 3 gagal → pulih: 15 menit.'),('40 menit','Studi Docker/React/Compose: 15 menit.\nDeployment: 15 menit.\nEvaluasi: 10 menit.')],note='Target: menjelaskan event/job/step/runner, menjalankan manual workflow, membaca log gagal, dan menghubungkan CI dengan proses deployment. Total 120 menit, persiapan akun dilakukan sebelumnya.')
cards('Istilah','Baca workflow dari besar ke kecil',[('Workflow + event','Workflow: satu file otomasi.\nEvent: pemicunya, misalnya push atau pull request.'),('Job + runner','Job: kelompok pekerjaan.\nRunner: mesin yang menjalankan job tersebut.'),('Step + action','Step: satu langkah.\nAction: komponen siap pakai yang dipanggil oleh uses.')],note='File workflow berada di .github/workflows dengan ekstensi .yml atau .yaml. Secara default step biasa berurutan; job tanpa needs dapat berjalan paralel sesuai kapasitas.')
s=slide('Alur kerja','Dari perubahan kode ke hasil pemeriksaan','Contoh: pull request memicu satu job CI.')
for i,(h,b) in enumerate([('Event','Pull request'),('Workflow','ci-node.yml'),('Runner','Checkout → test'),('Hasil','Lulus / gagal')]):
 x=.6+i*3.1;box(s,x,2.65,2.7,2.1,PANEL);text(s,x+.2,2.95,2.3,.5,h,24,ACCENT, bold=True);text(s,x+.2,3.7,2.3,.7,b,18)
 if i<3:text(s,x+2.78,3.3,.3,.4,'→',21,ACCENT)
text(s,.7,5.25,11.8,.8,'Buka run → pilih job → buka step → baca log.',25)
notes(s,'Tidak semua kegagalan berasal dari kode: bisa permission, jaringan, dependency, atau konfigurasi YAML. Mulai investigasi dari step pertama yang gagal.')
cards('Persiapan','Siapkan tempat belajar yang terpisah',[('Akun & repo','Akun GitHub dengan akses tulis.\nRepository latihan sendiri.\nGitHub Actions diizinkan.'),('File latihan','node-demo/server.js\nDua workflow contoh.\nScript smoke-node.mjs.'),('Batas latihan','Lab awal tanpa secret dan server deploy.\nJangan menyalin workflow deployment aktif.')],note='Repo sumber memiliki .github/workflows/deploy-react-demo.yml dengan trigger push main. Gunakan repo baru yang hanya memuat file latihan pada panduan. Policy organisasi dan kuota Actions dapat membatasi eksekusi.')
code('YAML','Indentasi menunjukkan struktur','name: Contoh\non:\n  workflow_dispatch:\njobs:\n  hello:\n    runs-on: ubuntu-latest\n    steps:\n      - name: Sapa peserta\n        run: echo "Halo"','jobs berisi job.\nsteps berisi daftar.\nTanda - memulai item.\nGunakan spasi, bukan tab.',note='YAML ini contoh minimal utuh. Perhatikan hello adalah ID job; Sapa peserta adalah nama step. Workflow lebih lengkap di sample menambahkan permissions dan timeout.')
code('Lab 1 • 15 menit','Workflow pertama: Hello Actions','name: Lab 1 - Hello Actions\non:\n  workflow_dispatch:\npermissions:\n  contents: read\njobs:\n  hello:\n    runs-on: ubuntu-latest\n    steps:\n      - run: echo "Halo dari GitHub Actions"','Simpan sebagai:\n.github/workflows/\nhello-actions.yml\n\nCommit ke default branch.',note='Gunakan sample lengkap samples/github-actions/hello-actions.yml. Slide menampilkan versi ringkas utuh; sample menambahkan timeout, checkout, dan ls. Default branch biasanya main.')
cards('Lab 1 • Verifikasi','Jalankan dan baca hasilnya',[('01 / Pilih','Buka tab Actions.\nPilih Lab 1 - Hello Actions.\nKlik Run workflow.'),('02 / Baca','Pilih run yang baru.\nBuka job hello.\nBuka step Sapa peserta.'),('03 / Buktikan','Log berisi pesan Halo.\nJob berstatus sukses.\nCatat commit yang dijalankan.')],note='Tombol Run workflow membutuhkan workflow_dispatch dan workflow yang ada di default branch; akses tulis diperlukan. https://docs.github.com/en/actions/how-tos/manage-workflow-runs/manually-run-a-workflow')
code('Step','uses memakai action; run menjalankan perintah','steps:\n  - uses: actions/checkout@v6\n    with:\n      persist-credentials: false\n  - name: Lihat source\n    run: ls -la','Checkout mengambil source ke workspace runner.\n\nTanpa checkout, jangan menganggap file repo sudah ada.',sub='Potongan di dalam job; bukan workflow lengkap.',note='Tag v6 dipilih agar konsisten dengan repo dan telah dicek pada dokumentasi tag resmi. Tidak disebut sebagai versi paling baru. Runner hosted dikelola GitHub; self-hosted perlu memenuhi versi minimum action.')
code('Event','Kapan pemeriksaan dijalankan?','on:\n  workflow_dispatch:\n  pull_request:\n  push:\n    branches: [main]','Manual: tombol Run.\nPR: usulan perubahan.\nPush main: commit masuk ke main.',sub='Trigger dari sample ci-node-pemula.yml.',note='Tanpa branches filter pada pull_request, workflow merespons PR ke branch target mana pun sesuai jenis event default. Sesuaikan main jika nama default branch berbeda. Fork PR mungkin perlu persetujuan kebijakan GitHub.')
cards('Job & runner','Job bisa paralel; needs membuat urutan',[('Tanpa needs','Job A dan B bisa berjalan bersamaan.\nStep biasa di satu job berjalan berurutan.'),('Dengan needs','Job deploy menunggu build berhasil.\nneeds: build'),('Data antar-job','Jangan menganggap file lokal tersedia di job lain.\nPindahkan lewat artifact, output, atau registry.')],note='GitHub-hosted VM umumnya baru untuk setiap job; state self-hosted berbeda. needs mengatur ketergantungan, bukan menyalin file antar runner.')
code('Lab 2 • 25 menit','Siapkan Node.js lalu cek sintaks','steps:\n  - uses: actions/checkout@v6\n  - uses: actions/setup-node@v6\n    with:\n      node-version: "22"\n      package-manager-cache: false\n  - name: Periksa sintaks\n    run: node --check node-demo/server.js','Source di-checkout.\nVersi Node dipilih.\nSintaks dicek sebelum aplikasi dijalankan.',sub='Potongan job dari ci-node-pemula.yml.',note='Sample lengkap punya permissions contents: read, timeout 5 menit, dan concurrency per ref. Node demo tanpa dependency eksternal sehingga tidak perlu npm ci.')
code('Lab 2 • HTTP test','Aplikasi bisa start belum berarti benar','- name: Uji HTTP dan isi respons\n  run: node samples/github-actions/smoke-node.mjs','Test memeriksa:\n/health → 200 + ok\n/api/info → JSON valid\n/tidak-ada → 404',sub='Script menjalankan server sementara dan menghentikannya setelah tes.',note='Script memilih port kosong, menunggu readiness hingga 10 detik, memeriksa field JSON, lalu membersihkan proses dalam finally. Ini smoke test, belum pengganti semua unit/integration test. Tidak perlu Docker atau secret.')
cards('Lab 2 • Verifikasi','Buktikan CI mengikuti perubahan',[('Pasang','Salin CI dan script ke path pada panduan.\nCommit ke repo latihan.'),('Jalankan','Run workflow manual.\nBuat branch lalu PR.\nAmati check di PR.'),('Hasil','Syntax dan HTTP hijau.\nLog menampilkan PASS.\nCommit run cocok dengan perubahan.')],note='Sample berada di samples sehingga belum aktif. Peserta menyalin ke .github/workflows pada repo latihan. Branch protection diperlukan bila CI harus wajib lulus sebelum merge; check hijau sendiri tidak mengatur proteksi.')
cards('Lab 3 • 15 menit','Sengaja gagal, lalu pulihkan',[('Rusakkan','Di branch latihan, ganti status /health dari ok ke rusak.\nCommit dan push.'),('Temukan','Syntax tetap lulus.\nHTTP test gagal.\nBaca expected dan actual pada log.'),('Perbaiki','Kembalikan nilai ok.\nPush commit perbaikan.\nRun terbaru harus hijau.')],note='Perubahan hanya pada repo latihan dan branch percobaan. Jangan langsung merge eksperimen rusak. Tes ini menunjukkan pemeriksaan konten lebih kuat daripada hanya memeriksa proses hidup.')
code('Dari Docker ke CI','Jalankan perintah Docker yang sudah dikenal','docker build -t docker-demo:ci ./docker-demo\ndocker run -d --name web-ci \\\n  -p 127.0.0.1:8080:80 docker-demo:ci\ncurl --fail --retry 10 --retry-connrefused \\\n  --retry-delay 1 http://127.0.0.1:8080','Build image.\nJalankan container.\nUji respons HTTP.\nTambahkan pemeriksaan isi.',sub='Ilustrasi run step pada runner Ubuntu; bukan workflow lengkap.',note='Tambahkan logs saat gagal dan cleanup if: always() untuk web-ci serta image tes. Jangan memakai global prune pada shared/self-hosted runner. Build sukses saja belum membuktikan aplikasi berjalan.')
code('Studi React','Pastikan perintah berjalan di folder aplikasi','- uses: actions/setup-node@v6\n  with:\n    node-version: "22"\n    cache: npm\n    cache-dependency-path: react-demo/package-lock.json\n- run: npm ci\n  working-directory: react-demo\n- run: npm run build\n  working-directory: react-demo','Checkout dilakukan sebelumnya.\n\nLockfile harus cocok.\nBuild menghasilkan dist/.',sub='Potongan steps; working-directory berlaku untuk step run.',note='Materi panjang sebelumnya menampilkan npm ci tanpa working-directory; jelaskan path di repo multi-aplikasi. Cache setup-node menyimpan data package manager, bukan node_modules. Detail versi tersedia di github.com/actions/setup-node/tree/v6.')
code('Studi Compose','Uji aplikasi bersama database','- name: Jalankan stack\n  run: docker compose up -d --build --wait\n  working-directory: compose-demo\n- name: Periksa layanan\n  run: curl --fail http://127.0.0.1:3001/health\n- name: Cleanup data tes\n  if: always()\n  run: docker compose down --volumes\n  working-directory: compose-demo','Gunakan runner tes sementara.\n\n--volumes menghapus data!\nJangan arahkan ke production.',sub='Potongan steps setelah checkout; tambah log saat gagal.',note='Ini perlu ditambah request root dan pemeriksaan SQL bila hendak membuktikan counter menulis data. /health hanya SELECT 1. Compose demo menggunakan PostgreSQL, bukan MySQL.')
cards('Hasil build','Cache dan artifact punya tujuan berbeda',[('Cache','Mempercepat run berikutnya.\nContoh: cache unduhan npm.\nPipeline harus benar saat cache kosong.'),('Artifact','Menyimpan hasil suatu run.\nContoh: laporan test atau dist.\nBisa diunduh atau dipakai job berikutnya.')],note='Artifact bukan deployment otomatis. File image Docker biasanya diterbitkan ke registry, bukan dianggap ikut ke runner job selanjutnya. https://docs.github.com/en/actions/concepts/workflows-and-actions/workflow-artifacts')
code('Konfigurasi','Pisahkan konfigurasi dan credential','env:\n  APP_ENV: latihan\n\n# Contoh input action yang memerlukan login:\nwith:\n  username: ${{ vars.DOCKER_USERNAME }}\n  password: ${{ secrets.DOCKERHUB_TOKEN }}','env: nilai di workflow.\nvars: konfigurasi GitHub.\nsecrets: credential.\n\nJangan cetak secret ke log.',sub='Dua potongan berbeda, bukan satu workflow utuh.',note='Lab 1–3 tidak membutuhkan secret. Untuk deployment sample repository variable bernama DOCKER_USERNAME; workflow aktif memakai secrets.DOCKER_USERNAME. Pahami masing-masing file sebelum menyiapkan konfigurasi.')
cards('Kontrol workflow','Mulai dengan izin dan waktu yang cukup',[('Permissions','CI: contents: read.\nTambahkan izin hanya sesuai pekerjaan.\nPin SHA action untuk produksi.'),('Timeout','Batasi durasi job.\nContoh: 5 menit untuk CI sederhana.\nLog harus membantu diagnosis.'),('Concurrency','CI: run lama boleh dibatalkan.\nDeploy: hindari memutus proses penggantian versi.')],note='Jangan menjalankan kode PR tidak tepercaya bersama credential produksi. Self-hosted runner memerlukan isolasi dan cleanup lebih ketat. Tag mayor digunakan agar contoh kelas mudah dibaca; tag dapat berpindah.')
s=slide('Gambaran CD','CI lulus → rilis → deploy → verifikasi','Deployment adalah materi lanjutan setelah CI dipahami.')
for i,(h,b) in enumerate([('Validasi','Test + build'),('Publikasi','Image ke registry'),('Deploy','Server menarik image'),('Verifikasi','Cek layanan')]):
 x=.6+i*3.1;box(s,x,2.65,2.7,2.1,PANEL);text(s,x+.18,2.95,2.34,.5,h,23,ACCENT, bold=True);text(s,x+.18,3.7,2.34,.7,b,18)
 if i<3:text(s,x+2.78,3.3,.3,.4,'→',21,ACCENT)
text(s,.7,5.25,11.8,.8,'Health gagal? Pulihkan versi lama dan laporkan deployment gagal.',23)
notes(s,'Diagram ini rancangan pembelajaran, bukan klaim bahwa workflow deployment aktif sudah memiliki job CI terpisah. Workflow aktif saat ini build-and-push lalu deploy.')
cards('Persiapan CD','Sebelum menekan tombol deploy',[('Tujuan','Registry dan server lab.\nPort aplikasi tersedia.\nRuntime Docker siap.'),('Akses','Token registry.\nSSH key khusus deploy.\nHost fingerprint tepercaya.'),('Kontrol','Environment production.\nBatas branch + reviewer bila tersedia.\nHealth check dan rencana rollback.')],note='Environment bernama production saja tidak otomatis mengaktifkan approval. Required reviewer dan fitur proteksi bergantung paket serta visibilitas repo. Jangan gunakan contoh deploy ini di lab pemula.')
cards('Versi & rollback','Kenali versi yang sedang berjalan',[('Tag commit','sha-<commit> membantu penelusuran.\nTag masih dapat ditimpa.\nDigest mengikat isi image.'),('Health check','Cek respons aplikasi setelah deploy.\nProses running saja belum cukup.'),('Rollback','Simpan referensi versi sebelumnya.\nUji layanan setelah pemulihan.\nPergantian satu container bisa downtime.')],note='Tag commit bukan immutable secara teknis. Gagal pada pull/run/deploy juga perlu ditangani, bukan hanya HTTP health yang gagal. Jangan menjanjikan zero downtime atau rollback teruji dari contoh saja.')
cards('Peta repository','Bedakan bahan latihan dan workflow aktif',[('Sample baru','hello-actions.yml\nci-node-pemula.yml\nDi samples/github-actions/; belum aktif.'),('Sample deploy','deploy-react-demo.yml di samples.\nTrigger manual.\nVariable DOCKER_USERNAME.'),('Workflow aktif','Di .github/workflows/.\nDeploy React saat push main atau manual.\nKonfigurasi secret berbeda.')],note='Jangan menimpa workflow deployment aktif untuk latihan. Panduan praktikum menggunakan repo terpisah. Materi github-actions.md adalah referensi tambahan; beberapa potongannya perlu dirangkai sebelum dipakai.')
cards('Troubleshooting','Cari step pertama yang gagal',[('Workflow tak muncul','Cek .github/workflows/.\nCek workflow_dispatch.\nCek default branch dan akses.'),('Run merah','Baca log step gagal.\nCek path, sintaks, versi runtime, dan dependency.'),('Tertunda / skipped','Cek needs dan if.\nCek izin Actions, approval, kapasitas, dan kuota.')],note='Jangan hanya mengulang run tanpa memperbaiki penyebab. Kegagalan jaringan sesaat berbeda dengan assertion yang pasti gagal.')
cards('Evaluasi','Peserta siap lanjut jika bisa…',[('Menjelaskan','Apa pemicu workflow?\nApa beda job dan step?\nMengapa perlu checkout?'),('Menunjukkan','Run Hello sukses.\nCI otomatis pada PR.\nLog assertion yang gagal.'),('Memperbaiki','Kembalikan CI ke hijau.\nTemukan file/workdir salah.\nBedakan CI dengan deploy.')],note='Jawaban ringkas: event memicu; job berjalan di runner dan berisi step; checkout mengambil source; CI memeriksa, deploy menerapkan versi ke environment. CI hijau hanya mencakup pemeriksaan yang ditulis.')
s=slide('Referensi & bahan praktik','Lanjutkan dengan latihan kecil yang bisa dibuktikan')
text(s,.7,2.25,11.8,1.1,'praktik-github-actions.md\nsamples/github-actions/hello-actions.yml  ·  ci-node-pemula.yml',22,ACCENT)
text(s,.7,3.8,11.8,2.35,'docs.github.com/en/actions/get-started/understand-github-actions\ngithub.com/actions/checkout/tree/v6\ngithub.com/actions/setup-node/tree/v6\nPanduan resmi workflow manual dan artifact: lihat catatan slide.',18)
notes(s,'Referensi: https://docs.github.com/en/actions/get-started/understand-github-actions ; https://docs.github.com/en/actions/how-tos/manage-workflow-runs/manually-run-a-workflow ; https://github.com/actions/checkout/tree/v6 ; https://github.com/actions/setup-node/tree/v6 ; https://docs.github.com/en/actions/concepts/workflows-and-actions/workflow-artifacts ; https://docs.github.com/en/actions/concepts/workflows-and-actions/deployment-environments . Identitas: Muhammad Asdar — Distritek DevOps — asdardevs@gmail.com.')
r.core_properties.title='GitHub Actions — Dari Pemula ke CI/CD'
r.core_properties.author='Muhammad Asdar'
r.core_properties.subject='Workshop GitHub Actions, workflow, CI/CD, praktik dan deployment'
r.save(ROOT/'materi-github-actions.pptx')
print(f'Generated {len(r.slides)} slides')
