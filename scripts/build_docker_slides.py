"""Generate the revised workshop deck: python -m pip install python-pptx==1.0.2."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
r = Presentation()
r.slide_width, r.slide_height = Inches(13.333), Inches(7.5)
BG, PANEL, INK, MUTED, ACCENT = '101D30', '1B2C43', 'F4F7FB', 'B5C4D9', '45DECB'

def box(s,x,y,w,h,color):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
    return sh

def text(s,x,y,w,h,content,size=23,color=INK,font='Aptos',bold=False):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02); tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12 if font!='Consolas' else 3)
        if font=='Consolas': p.line_spacing=1.0
    return sh

def slide(section,title,sub=None):
    s=r.slides.add_slide(r.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    box(s,0,0,.16,7.5,ACCENT)
    text(s,.6,.35,12,.3,section.upper(),12,ACCENT,bold=True)
    text(s,.6,.88,12.1,.65,title,32,bold=True)
    if sub: text(s,.6,1.62,12,.5,sub,17,MUTED)
    box(s,.6,6.98,12.1,.012,'38516D')
    text(s,.6,7.13,11,.22,'DISTRITEK DEVOPS   /   Muhammad Asdar   /   Docker Workshop',10,MUTED)
    text(s,12.05,7.08,.6,.3,str(len(r.slides)).zfill(2),12,ACCENT)
    return s

def notes(s,body): s.notes_slide.notes_text_frame.text=body

def cards(section,title,items,sub=None,note=''):
    s=slide(section,title,sub); n=len(items); w=(12.1-.25*(n-1))/n
    for i,(head,body) in enumerate(items):
        x=.6+i*(w+.25); box(s,x,2.3,w,3.9,PANEL)
        text(s,x+.22,2.57,w-.44,.8,head,24,ACCENT,bold=True)
        text(s,x+.22,3.5,w-.44,2.5,body,21)
    notes(s,note); return s

def code(section,title,body,explanation,sub=None,note=''):
    s=slide(section,title,sub); box(s,.6,2.25,8,4.38,PANEL)
    longest=max(map(len,body.splitlines())); count=len(body.splitlines())
    size=min(20, 890/max(longest,1), 278/max(count,1)-3)
    text(s,.85,2.47,7.5,3.95,body,size,INK,'Consolas')
    text(s,8.98,2.4,3.7,3.9,explanation,22)
    notes(s,note); return s

s=slide('Workshop praktik • 150 menit','Docker: dari image ke aplikasi nyata')
text(s,.7,2.2,8,1.7,'Build sekali.\nPahami cara menjalankannya.',32,bold=True)
text(s,.7,4.25,7.5,1.3,'4 demo bertahap: Nginx, Node.js, React,\ndan Node.js + PostgreSQL.',24,MUTED)
box(s,9,2.35,3.65,3.5,PANEL); text(s,9.3,2.7,3,2.7,'BUILD\n↓\nRUN\n↓\nVERIFY',27,ACCENT,bold=True)
text(s,.7,6.1,9,.4,'Muhammad Asdar  ·  Distritek DevOps',19)
notes(s,'Materi pemula dengan asumsi peserta mengenal terminal dasar. Instalasi dilakukan sebelum kelas. Panduan lengkap dan perintah copy-paste tersedia di praktik-docker.md.')
cards('Perkenalan','Tentang fasilitator',[('Muhammad Asdar','DevOps & Software Engineer\nDistritek DevOps'),('Terhubung','Telegram: @asdarmld\nLinkedIn: /in/asdarmld\nInstagram: @asdarmld')],note='Identitas dipertahankan dari presentasi asli. Email: asdardevs@gmail.com')
cards('Tujuan belajar','Selesai kelas, peserta bisa…',[('Jalankan','Bedakan image dan container.\nBuild image, atur port, baca log.'),('Jelaskan','Bedakan COPY dan mount.\nTelusuri jaringan dan data.'),('Buktikan','Jalankan 4 demo.\nTunjukkan database tetap ada setelah container diganti.')])
cards('Rute kelas','Konsep singkat, praktik bertahap',[('01 / Fondasi','Konsep: 25 menit\nPersiapan: 10 menit\nNginx: 15 menit'),('02 / Build & run','Node.js: 20 menit\nReact: 20 menit\nCompose: 30 menit'),('03 / Evaluasi','Diskusi, troubleshooting, dan cleanup: 30 menit\nTotal: 150 menit')],note='Unduhan image awal tidak termasuk durasi. Jika waktu terbatas, React dapat dipindahkan menjadi tugas mandiri.')
cards('Pengantar • Mulai dari masalah','Mengapa aplikasi kadang hanya jalan di satu laptop?',[('Di laptop pembuat','Aplikasi memakai Node.js versi tertentu dan library yang sudah terpasang.\nSemua berjalan normal.'),('Di laptop teman','Versi atau library berbeda.\nAplikasi yang sama bisa gagal dijalankan.')],sub='Kode yang sama masih membutuhkan lingkungan yang sesuai.',note='Mulai dengan pertanyaan: pernah mengirim project ke teman, tetapi muncul error saat dijalankan? Jelaskan dependensi sebagai bahan pendukung yang diperlukan aplikasi. Jangan langsung memakai istilah daemon, registry, atau kernel.')
cards('Pengantar • Ide dasarnya','Docker membantu menyiapkan paket aplikasi',[('Paketnya','Kode aplikasi + runtime + library yang diperlukan.\nRuntime adalah program yang menjalankan kode, misalnya Node.js.'),('Cara memakainya','Docker menjalankan paket tersebut dalam lingkungan proses yang terisolasi.\nKita tidak perlu menyiapkan semuanya satu per satu.')],sub='Tujuannya: mengurangi perbedaan lingkungan saat aplikasi dipindahkan.',note='Naskah pembuka: Docker adalah alat untuk mengemas dan menjalankan aplikasi beserta kebutuhan utamanya. Bayangkan kita berbagi paket aplikasi, bukan hanya mengirim kode lalu meminta teman menebak apa yang harus diinstal. Tetap diperlukan Docker, platform yang kompatibel, dan konfigurasi runtime yang benar.')
cards('Pengantar • Analogi sederhana','Resep → paket → aplikasi berjalan',[('Dockerfile = resep','Catatan tentang apa yang dibutuhkan dan cara menyiapkan paket aplikasi.'),('Image = paket','Hasil jadi yang dapat disimpan dan dipakai berulang kali.\nBelum berarti aplikasi sedang berjalan.'),('Container = berjalan','Satu lingkungan tempat paket dijalankan.\nSatu image bisa menghasilkan beberapa container.')],sub='Ini analogi untuk membantu mengenal istilah, bukan detail cara kerja.',note='Tunjuk diagram tiga tahap. Contoh kita akan membuat paket halaman web, lalu menjalankannya. Analogi tidak berarti image berisi kernel lengkap. Container juga tetap ada ketika prosesnya stopped.')
cards('Konsep','Image ≠ container',[('Image','Paket read-only berisi aplikasi dan dependensi runtime.\nSatu image dapat dipakai banyak container.'),('Container','Instance image dengan proses dan writable layer sendiri.\nBisa running atau stopped.')],sub='Docker membantu konsistensi environment; kompatibilitas platform tetap penting.',note='Hindari janji hasil identik di semua mesin: arsitektur CPU, kernel, konfigurasi runtime, mount, dan layanan eksternal tetap berpengaruh.')
s=slide('Konsep','Alur kerja Docker','Dockerfile membangun image; container menjalankan aplikasi.')
for i,(h,b) in enumerate([('Dockerfile','Resep build'),('Image','Hasil build'),('Container','Instance runtime'),('Registry','Push / pull image')]):
 x=.6+i*3.1; box(s,x,2.6,2.7,2.05,PANEL); text(s,x+.18,2.9,2.35,.5,h,25,ACCENT, bold=True); text(s,x+.18,3.65,2.35,.5,b,19)
 if i<2: text(s,x+2.77,3.2,.3,.5,'→',23,ACCENT)
text(s,.7,5.1,11.8,.9,'Image lokal bisa langsung di-run. Registry dipakai untuk distribusi.',24)
notes(s,'Registry tidak wajib untuk run image lokal. Diagram menunjukkan siklus distribusi; image lokal dapat langsung dijalankan tanpa push/pull. CLI mengirim permintaan ke Engine/daemon.')
cards('Konsep','Container dan VM: beda lapisan isolasi',[('Container','Isolasi proses; berbagi kernel lingkungan tempat berjalan.\nUmumnya startup dan overhead lebih ringan.'),('Virtual machine','OS dan kernel sendiri di atas hypervisor.\nOverhead umumnya lebih besar.')],sub='Ukuran dan waktu startup bergantung workload, bukan angka tetap.',note='Pada Docker Desktop, Linux container berjalan di VM Linux. Image membawa user-space OS, bukan kernel lengkap. Container bukan batas keamanan yang identik dengan VM.')
code('Persiapan','Pastikan engine benar-benar siap','docker version\ndocker compose version\ndocker run --rm hello-world','Lulus:\nClient + Server terdeteksi.\nCompose tersedia.\nHello-world berhasil.',sub='Linux: Engine + Compose plugin • macOS / Windows: Docker Desktop',note='Rujuk panduan instalasi resmi di praktik-docker.md. Windows gunakan WSL2. Grup docker memberi hak setingkat root. Podman tidak diasumsikan identik dan perlu Compose provider tersendiri.')
code('Lifecycle','Run membuat container baru','docker run -d --name lab-web \\\n  -p 127.0.0.1:8080:80 nginx:alpine\ndocker ps\ndocker logs lab-web\ndocker stop lab-web\ndocker start lab-web\ndocker stop lab-web\ndocker rm lab-web','run: buat + mulai\nstop: hentikan\nstart: container lama\nrm: hapus container',note='Jalankan contoh lifecycle sampai cleanup sebelum lab 1 agar nama lab-web tidak bentrok. Ctrl+C keluar dari docker logs -f tanpa menghentikan container.')
cards('Jaringan','Baca port dari kiri ke kanan',[('127.0.0.1:8080','Alamat dan port di host.\nBuka localhost:8080 di browser.'),('80','Port di container.\nHarus cocok dengan port listen aplikasi.')],sub='-p 127.0.0.1:8080:80',note='Binding localhost membatasi akses ke host. EXPOSE tidak memublikasikan port. Untuk host remote gunakan SSH tunnel. localhost dalam container menunjuk dirinya sendiri.')
code('Lab 1 • 15 menit','Build halaman statis Anda','docker build -t docker-demo:latihan ./docker-demo\ndocker run -d --name lab-web \\\n  -p 127.0.0.1:8080:80 docker-demo:latihan\ncurl -f http://localhost:8080\ndocker exec lab-web nginx -t','Lulus:\nHTML tampil di 8080.\nNginx config valid.\n\nSumber: docker-demo/',sub='Jalankan dari root repository.',note='Baca Dockerfile tiga instruksi. Ubah judul index.html, refresh, lalu minta peserta menjelaskan mengapa halaman belum berubah. Build ulang dan buat ulang container mengikuti panduan.')
code('Build','Dockerfile: build time vs runtime','FROM node:22-alpine\nWORKDIR /app\nCOPY package.json ./\nCOPY server.js ./\nUSER node\nEXPOSE 3000\nCMD ["node", "server.js"]','COPY: saat build\nCMD: saat run\nUSER: non-root\nEXPOSE: dokumentasi',sub='Contoh node-demo tanpa dependensi eksternal; tidak perlu npm install.',note='FROM memulai stage; ARG atau parser directive boleh mendahuluinya. WORKDIR mengatur direktori kerja. Titik terakhir docker build adalah build context. Aplikasi Node listen 0.0.0.0.')
code('Build','Cache dependensi + lockfile','COPY package*.json ./\nRUN npm ci --omit=dev\nCOPY server.js ./','package.json +\npackage-lock.json\nharus cocok.\n\nKode berubah → layer instalasi bisa dipakai ulang.',sub='Pola compose-demo/app • Untuk React, npm ci juga memasang build tools.',note='npm ci gagal bila lockfile tidak ada atau tidak cocok. Sertakan lockfile di Git. .dockerignore mengecualikan node_modules, .git, .env, secrets, dan dist untuk build React. https://docs.npmjs.com/cli/v11/commands/npm-ci/')
cards('Storage','COPY, bind mount, dan volume',[('COPY','Snapshot file ke image saat build.\nEdit source → build ulang + container baru.'),('Bind mount','File host dipasang ke container.\nCocok untuk edit cepat; gunakan read-only bila cukup.'),('Named volume','Penyimpanan dikelola Docker.\nCocok untuk data database. Volume bukan backup.')],note='Writable layer bertahan saat stop/start, hilang ketika container dihapus. Named volume tidak hilang oleh compose down biasa. Mount dapat menutupi isi path yang sudah ada di image.')
code('Lab 1 • Eksperimen','Edit tanpa build ulang','LAB_HTML="$(pwd)/docker-demo/index.html"\ndocker run -d --name lab-bind \\\n  -p 127.0.0.1:8082:80 \\\n  -v "$LAB_HTML:/usr/share/nginx/html/index.html:ro" \\\n  nginx:alpine','Lulus:\nEdit HTML → refresh 8082 langsung berubah.\n\nBandingkan dengan 8080.',sub='Perintah dijalankan dari root repository.',note='Windows gunakan WSL2. Bind mount read-only mencegah container mengubah file host.')
code('Lab 2 • 20 menit','Node.js: amati aplikasi dinamis','docker build -t node-demo:latihan ./node-demo\ndocker run -d --name lab-node \\\n  -p 127.0.0.1:3000:3000 node-demo:latihan\ncurl -f http://localhost:3000/api/info\ncurl -f http://localhost:3000/health\ndocker exec lab-node id\ndocker logs lab-node','Lulus:\nJSON info tampil.\nHealth: status ok.\nUser bukan root.',note='Endpoint root menyajikan halaman dinamis. Uptime adalah uptime proses Node; restart mengulang uptime. Hostname default biasanya ID container pendek, tetapi bisa dikonfigurasi.')
code('Lab 2 • Eksperimen','Satu image, dua identitas runtime','docker run -d --name lab-node-env \\\n  -e PORT=4000 \\\n  -p 127.0.0.1:3002:4000 node-demo:latihan\ncurl -f http://localhost:3002/api/info','Bandingkan hostname dengan port 3000.\n\nPORT=4000 mengubah port aplikasi di container.',note='Tantangan: port mapping kanan salah 3000 ketika app listen 4000. Identifikasi dengan logs. Jangan menggunakan port 3001 karena dipakai Compose.')
code('Lab 3 • 20 menit','React: pisahkan build dan runtime','FROM node:22-alpine AS builder\nWORKDIR /app\nCOPY package*.json ./\nRUN npm ci\nCOPY . .\nRUN npm run build\n\nFROM nginx:alpine\nCOPY nginx.conf /etc/nginx/conf.d/default.conf\nCOPY --from=builder /app/dist /usr/share/nginx/html','Stage 1: build aset.\nStage 2: Nginx + dist.\n\nNode.js tidak ikut ke image runtime.',sub='Ringkasan Dockerfile di react-demo/; ukuran image diukur, tidak diasumsikan.')
code('Lab 3 • Verifikasi','Pastikan route SPA bisa di-refresh','docker build -t react-demo:latihan ./react-demo\ndocker run -d --name lab-react \\\n  -p 127.0.0.1:8081:80 react-demo:latihan\ncurl -f http://localhost:8081/materi\ndocker exec lab-react sh -c \\\n  "command -v nginx; command -v node || true"','Buka /materi langsung di browser.\nRefresh tetap berhasil.\n\nNginx ada, Node tidak ada.',note='Fallback try_files ke /index.html diperlukan karena routing React terjadi di browser. curl memeriksa HTML fallback; verifikasi UI tetap melalui browser.')
cards('Network','Aplikasi menemukan database lewat DNS',[('Host → app','localhost:3001\nDipublikasikan ke port 3000 container app.'),('app → db','DB_HOST=db\nPostgreSQL port 5432 lewat network Compose.'),('db → volume','dbdata\nMount ke /var/lib/postgresql/data untuk PostgreSQL 16.')],sub='app dan db berada di network Compose yang sama.',note='Resolusi nama otomatis berlaku pada user-defined network, bukan default bridge dengan cara yang sama. Database tidak perlu ports untuk diakses app pada network yang sama.')
code('Compose • 1/2','Konfigurasi aplikasi','services:\n  app:\n    build: ./app\n    ports: ["127.0.0.1:3001:3000"]\n    environment:\n      DB_HOST: db\n      DB_USER: demo\n      DB_PASSWORD: rahasia\n      DB_NAME: demo\n    depends_on:\n      db:\n        condition: service_healthy','Potongan dari compose-demo/compose.yaml.\n\nFile lengkap tersedia di repo.\nPassword ini khusus lab.',note='Slide ini potongan, bukan Compose lengkap. File repo juga menyediakan healthcheck app. Health endpoint menjalankan SELECT 1 tanpa menambah kunjungan.')
code('Compose • 2/2','Database + healthcheck + volume','  db:\n    image: postgres:16-alpine\n    environment:\n      POSTGRES_USER: demo\n      POSTGRES_PASSWORD: rahasia\n      POSTGRES_DB: demo\n    volumes: ["dbdata:/var/lib/postgresql/data"]\n    healthcheck:\n      test: ["CMD-SHELL", "pg_isready -U demo -d demo"]\n      interval: 3s\n      timeout: 3s\n      retries: 10\nvolumes:\n  dbdata:','db menjorok di bawah services.\nvolumes di level root.\n\nJalankan file lengkap dari repo.',note='Kredensial demo tidak untuk produksi. Environment inisialisasi hanya diterapkan ketika data directory kosong. pg_isready mengukur readiness server, sedangkan health app memeriksa query menggunakan koneksi app.')
code('Lab 4 • 30 menit','Jalankan dua service bersama','cd compose-demo\ndocker compose config -q\ndocker compose up -d --build --wait\ndocker compose ps\ncurl -f http://localhost:3001/health\ncurl -f http://localhost:3001/\ndocker compose logs app','Lulus:\nApp + DB healthy.\nHalaman counter tampil.\nRefresh menambah kunjungan.',sub='Mulai bagian ini, perintah Compose dijalankan dari folder compose-demo.',note='--wait menunggu running/healthy; butuh Compose modern. Instalasi dilakukan di prasyarat. config -q memvalidasi tanpa mencetak environment.')
code('Lab 4 • Persistensi','Buktikan data tetap ada','docker compose exec db psql -U demo -d demo \\\n  -tAc "SELECT count(*) FROM kunjungan;"\ndocker compose down\ndocker compose up -d --wait\ndocker compose exec db psql -U demo -d demo \\\n  -tAc "SELECT count(*) FROM kunjungan;"','Catat nilai N.\n\nSetelah down/up:\nSQL tetap N sebelum halaman / dibuka lagi.',note='Healthcheck tidak menyentuh counter. Gunakan project dan file yang sama. down -v berbeda: menghapus named volume proyek.')
code('Lab 4 • Ketahanan','Running belum tentu sehat','docker compose stop db\ncurl -i http://localhost:3001/health\ndocker compose start db\n# Ulangi setelah DB siap:\ncurl -i http://localhost:3001/health','DB mati → HTTP 503\nDB siap → HTTP 200\n\nHealthcheck tidak otomatis me-restart app.',note='depends_on service_healthy mengatur startup, bukan jaminan kesehatan sepanjang runtime. App menangani error pool idle dan menggunakan timeout query. restart policy bereaksi pada proses yang berhenti. https://docs.docker.com/compose/how-tos/startup-order/')
code('Cleanup','Hapus hanya resource latihan','docker compose down\n# Dari direktori mana pun:\ndocker stop lab-web lab-bind lab-node \\\n  lab-node-env lab-react\ndocker rm lab-web lab-bind lab-node \\\n  lab-node-env lab-react','Volume tetap ada.\n\nReset data sengaja:\ndocker compose down -v\nData counter hilang.',note='Jika tidak semua lab dijalankan, pakai hanya nama container yang dibuat. Hindari system prune -a sebagai default kelas: dampaknya lintas proyek.')
cards('Troubleshooting','Periksa gejalanya, lalu buktikan',[('Tidak bisa akses','docker ps -a\ndocker logs NAMA\nCek port kiri/kanan dan listen 0.0.0.0.'),('Build gagal','Cek build context.\nCek package-lock.json.\nGunakan sh jika bash tidak ada.'),('Data / route salah','DB_HOST harus db.\nCek named volume.\nCek fallback SPA dan build ulang.')])
cards('Kebiasaan baik','Dari latihan menuju deployment',[('Image','Pakai tag terkontrol; digest untuk pin immutable.\nSertakan lockfile dan .dockerignore.'),('Runtime','Non-root bila didukung.\nKelola secret saat runtime.\nBatasi CPU/memori dan pantau health.'),('Operasional','Backup data volume.\nUji restore dan recovery.\nUpdate image secara berkala.')],note='Alpine memakai musl dan tidak selalu cocok untuk semua dependency. Tag dapat berubah, termasuk versi. Demo Nginx standar belum menerapkan non-root khusus. Lab ini bukan template produksi lengkap.')
cards('Evaluasi','Tunjukkan hasil, bukan hanya perintah',[('Jelaskan','Mengapa restart tidak memuat image baru?\nMengapa EXPOSE belum cukup?'),('Buktikan','Dua hostname Node berbeda.\n/materi bisa di-refresh.\nSQL tetap N setelah down/up.'),('Perbaiki','Port app salah.\nDB berhenti sementara.\nImage belum dibuild ulang.')],note='Jawaban: restart memakai container lama; perlu recreate untuk image baru. EXPOSE metadata; -p memublikasikan. Persistensi berasal dari named volume, bukan writable layer.')
s=slide('Referensi & tindak lanjut','Bahan latihan dan dokumentasi')
text(s,.7,2.25,11.8,1.2,'praktik-docker.md  →  panduan langkah demi langkah\ndocker-demo/  ·  node-demo/  ·  react-demo/  ·  compose-demo/',23,ACCENT)
text(s,.7,3.75,11.8,2.3,'docs.docker.com/reference/dockerfile/\ndocs.docker.com/compose/how-tos/startup-order/\ndocs.docker.com/engine/storage/volumes/\ndocs.npmjs.com/cli/v11/commands/npm-ci/',21)
notes(s,'Referensi resmi diperiksa saat revisi. Detail analisis dan batas validasi di analisis-materi-docker.md. Terima kasih. Muhammad Asdar — asdardevs@gmail.com.')
r.core_properties.title='Materi Dasar Docker — Workshop Praktik'
r.core_properties.author='Muhammad Asdar'
r.core_properties.subject='Docker, Dockerfile, networking, storage, multi-stage build, Compose'
r.save(ROOT/'materi-docker-revisi.pptx')
print(f'Generated {len(r.slides)} slides')
